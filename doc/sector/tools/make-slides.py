#!/usr/bin/env python3
"""Make the sector modelling presentation, a PowerPoint deck for users.

Usage: make-slides.py <output.pptx>

The content follows the guide in the parent folder: the introduction and the
worked examples make up the main part, the limitations get one slide, and
the technical documentation is a set of optional slides at the end. Needs
python-pptx and Pillow.
"""

import os
import re
import sys

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
IMAGES = os.path.join(HERE, os.pardir, "images")

NAVY = RGBColor(0x1F, 0x3B, 0x5A)
ACCENT = RGBColor(0xD9, 0x6C, 0x1E)
GREY = RGBColor(0x55, 0x55, 0x55)
LIGHT = RGBColor(0xF3, 0xF5, 0xF8)
CODE_BG = RGBColor(0xF1, 0xF1, 0xF1)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLUE_CARD = RGBColor(0xE3, 0xEE, 0xF8)
GREEN_CARD = RGBColor(0xE6, 0xF2, 0xE1)
ORANGE_CARD = RGBColor(0xFB, 0xEB, 0xD9)
PINK_CARD = RGBColor(0xF6, 0xE3, 0xEC)

FONT = "Calibri"
MONO = "Consolas"

W, H = Inches(13.333), Inches(7.5)
LEFT = Inches(0.6)
TOP = Inches(1.35)
FULL_W = W - 2 * LEFT
BODY_H = Inches(5.55)

prs = Presentation()
prs.slide_width, prs.slide_height = W, H
BLANK = prs.slide_layouts[6]
TOKEN = re.compile(r"(\*\*.+?\*\*|`.+?`|_\{.+?\}|\^\{.+?\}|\*.+?\*)")


def add_runs(paragraph, text, size, color=None, bold=False, font=FONT):
    """Add text with **bold**, *italic*, `code`, _{sub} and ^{sup} markup."""
    for part in TOKEN.split(text):
        if not part:
            continue
        run = paragraph.add_run()
        run.font.size = Pt(size)
        run.font.name = font
        run.font.bold = bold
        if color is not None:
            run.font.color.rgb = color
        if part.startswith("**"):
            run.text = part[2:-2]
            run.font.bold = True
        elif part.startswith("`"):
            run.text = part[1:-1]
            run.font.name = MONO
            run.font.size = Pt(size - 1)
        elif part.startswith("_{") or part.startswith("^{"):
            run.text = part[2:-1]
            run.font._rPr.set("baseline", "-25000" if part[0] == "_" else "30000")
        elif part.startswith("*"):
            run.text = part[1:-1]
            run.font.italic = True
        else:
            run.text = part


def set_bullet(paragraph, level, char="•"):
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set("marL", str(int(Inches(0.28 + 0.32 * level))))
    pPr.set("indent", str(int(-Inches(0.24))))
    clr = pPr.makeelement(qn("a:buClr"), {})
    srgb = clr.makeelement(qn("a:srgbClr"), {"val": "D96C1E" if level == 0 else "1F3B5A"})
    clr.append(srgb)
    pPr.append(clr)
    pPr.append(pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"}))
    pPr.append(pPr.makeelement(qn("a:buChar"), {"char": char if level == 0 else "–"}))


def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    return box, tf


def bullets(slide, items, x=LEFT, y=TOP, w=FULL_W, h=BODY_H, size=20, gap=6):
    """Bullet list. An item is a string, or (level, string); a string that
    starts with '#' is an unbulleted heading."""
    _, tf = textbox(slide, x, y, w, h)
    first = True
    for item in items:
        level, text = item if isinstance(item, tuple) else (0, item)
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_before = Pt(gap if level == 0 else 2)
        if text.startswith("#"):
            p.space_before = Pt(gap + 4)
            add_runs(p, text[1:].strip(), size, NAVY, bold=True)
        else:
            set_bullet(p, level)
            add_runs(p, text, size if level == 0 else size - 2, RGBColor(0x22, 0x22, 0x22))
    return tf


def code(slide, text, x, y, w, h, size=13):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = CODE_BG
    shape.line.color.rgb = RGBColor(0xD0, 0xD0, 0xD0)
    shape.shadow.inherit = False
    tf = shape.text_frame
    tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = tf.margin_right = Inches(0.15)
    tf.margin_top = tf.margin_bottom = Inches(0.1)
    for i, line in enumerate(text.strip("\n").split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        run = p.add_run()
        run.text = line if line else " "
        run.font.name = MONO
        run.font.size = Pt(size)
        run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
    return shape


def caption(slide, text, x, y, w, size=14, align=PP_ALIGN.CENTER):
    _, tf = textbox(slide, x, y, w, Inches(0.4))
    p = tf.paragraphs[0]
    p.alignment = align
    add_runs(p, text, size, GREY)


def picture(slide, name, x, y, w, h):
    """Place an image as large as fits in the box, centred."""
    path = os.path.join(IMAGES, name)
    iw, ih = Image.open(path).size
    scale = min(w / iw, h / ih)
    pw, ph = int(iw * scale), int(ih * scale)
    return slide.shapes.add_picture(path, x + (w - pw) // 2, y + (h - ph) // 2, pw, ph)


def table(slide, rows, x, y, w, widths, size=13, row_h=0.32, align=None):
    """Table with a header row. widths are relative; align is a string of
    'l' and 'r' per column."""
    nr, nc = len(rows), len(rows[0])
    shape = slide.shapes.add_table(nr, nc, x, y, w, Inches(row_h * nr))
    tbl = shape.table
    total = sum(widths)
    for j, wj in enumerate(widths):
        tbl.columns[j].width = int(w * wj / total)
    align = align or "l" * nc
    for i, row in enumerate(rows):
        tbl.rows[i].height = Inches(row_h)
        for j, value in enumerate(row):
            cell = tbl.cell(i, j)
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY if i == 0 else (LIGHT if i % 2 == 0 else WHITE)
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.RIGHT if align[j] == "r" and i > 0 else PP_ALIGN.LEFT
            add_runs(p, value, size, WHITE if i == 0 else RGBColor(0x22, 0x22, 0x22), bold=i == 0)
    return shape


def card(slide, x, y, w, h, heading, items, fill, size=17):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    shape.adjustments[0] = 0.04
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    shape.shadow.inherit = False
    _, tf = textbox(slide, x + Inches(0.2), y + Inches(0.12), w - Inches(0.4), Inches(0.5))
    add_runs(tf.paragraphs[0], heading, size + 3, NAVY, bold=True)
    bullets(slide, items, x + Inches(0.15), y + Inches(0.62), w - Inches(0.3),
            h - Inches(0.7), size=size, gap=4)


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = " ".join(text.split())


def new_slide(title, optional=False):
    slide = prs.slides.add_slide(BLANK)
    _, tf = textbox(slide, LEFT, Inches(0.35), FULL_W - Inches(2.2), Inches(0.8),
                    anchor=MSO_ANCHOR.MIDDLE)
    add_runs(tf.paragraphs[0], title, 30, NAVY, bold=True)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, LEFT, Inches(1.15), Inches(1.2), Inches(0.06))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()
    if optional:
        _, tf = textbox(slide, W - LEFT - Inches(2.6), Inches(0.45), Inches(2.6), Inches(0.4))
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.RIGHT
        add_runs(p, "OPTIONAL · TECHNICAL", 12, ACCENT, bold=True)
    _, tf = textbox(slide, LEFT, H - Inches(0.45), Inches(6), Inches(0.3))
    add_runs(tf.paragraphs[0], "Sector modelling in OPM Flow", 11, GREY)
    _, tf = textbox(slide, W - LEFT - Inches(1), H - Inches(0.45), Inches(1), Inches(0.3))
    tf.paragraphs[0].alignment = PP_ALIGN.RIGHT
    add_runs(tf.paragraphs[0], str(len(prs.slides)), 11, GREY)
    return slide


def dark_slide(title, subtitle, image=None):
    slide = prs.slides.add_slide(BLANK)
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = NAVY
    bg.line.fill.background()
    text_w = Inches(6.6) if image else W - 2 * LEFT
    _, tf = textbox(slide, Inches(0.8), Inches(2.3), text_w, Inches(1.6), anchor=MSO_ANCHOR.BOTTOM)
    add_runs(tf.paragraphs[0], title, 40, WHITE, bold=True)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.85), Inches(4.05), Inches(1.6), Inches(0.08))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()
    _, tf = textbox(slide, Inches(0.8), Inches(4.3), text_w, Inches(2))
    for i, line in enumerate(subtitle.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(6)
        add_runs(p, line, 22 if i == 0 else 16, RGBColor(0xDD, 0xE6, 0xF0))
    if image:
        picture(slide, image, Inches(7.5), Inches(1.4), Inches(5.4), Inches(4.7))
    return slide


def two_columns(gap=Inches(0.4)):
    cw = (FULL_W - gap) // 2
    return cw, LEFT, LEFT + cw + gap


# ---------------------------------------------------------------- Introduction

s = dark_slide("Sector modelling\nin OPM Flow",
               "DUMPFLUX, USEFLUX and make_flux\nA user's guide, with worked examples on the Drogon model",
               image="usesector_pressure.png")
notes(s, """This presentation follows the sector modelling guide in
opm-simulators, doc/sector. The main part is for users: what the features do
and how to use them, with three worked examples. Technical details are in
optional slides at the end.""")

s = new_slide("Why sector modelling?")
bullets(s, [
    "Simulate **part of a field model**, a *sector*, with the rest of the field "
    "replaced by what crosses the sector's boundary",
    "The boundary data comes from a run of the full model, the **parent**, and "
    "is stored in a **FLUX file**",
    "The **reduced run** reproduces the parent inside the sector to within a "
    "few tenths of a percent …",
    "… at a fraction of the cost: **27 s instead of 105 s** on Drogon, "
    "where the sector is a quarter of the model",
    "Study changes inside the sector, such as a new well, a different "
    "completion or modified rock properties, without rerunning the field",
    "Eclipse keywords DUMPFLUX, USEFLUX, FLUXNUM, FLUXREG and FLUXTYPE, "
    "with extensions for wells, groups and UDQs outside the sector",
], w=Inches(6.9))
picture(s, "full_fluxnum.png", Inches(7.7), TOP + Inches(0.4), Inches(5.0), Inches(3.6))
caption(s, "Drogon with the sector in red (FLUXNUM 1)", Inches(7.7), TOP + Inches(4.1), Inches(5.0))
notes(s, """The idea: run the full field once, record what crosses the
boundary of the area of interest, and then iterate on the sector alone. The
reduced run is about four times faster on Drogon, in proportion to the number
of cells.""")

s = new_slide("A sector study in three steps")
picture(s, "workflow.png", LEFT, TOP, FULL_W, Inches(3.5))
cw = (FULL_W - Inches(0.6)) // 3
steps = [
    ("1  Define the sector", ["One region of **FLUXNUM**, over the parent's whole grid",
                              "EQUALS, BOX or INCLUDE"], BLUE_CARD),
    ("2  Produce the boundary", ["**DUMPFLUX** in the parent run, or",
                                 "**make_flux** on the output of an ordinary run"], ORANGE_CARD),
    ("3  Run the sector", ["**USEFLUX** + **FLUXREG**",
                           "On the parent deck, or on a cut-out sector deck"], PINK_CARD),
]
for i, (head, items, fill) in enumerate(steps):
    card(s, LEFT + i * (cw + Inches(0.3)), TOP + Inches(3.65), cw, Inches(1.85), head, items, fill, size=16)
notes(s, """The FLUX file is the hand-over between the parent and the reduced
run. It holds the boundary data, and also the parent's summary vectors and
wells, which the reduced run needs for everything outside the sector.""")

s = new_slide("Keywords")
table(s, [
    ["Keyword", "Section", "Purpose"],
    ["FLUXNUM", "GRID", "Region array defining the sectors over the whole grid; 0 means no sector"],
    ["DUMPFLUX", "GRID", "Write a FLUX file for each FLUXNUM region during the run"],
    ["FLUXTYPE", "GRID", "`'PRESSURE'`, `'FLUX'` or `'BOTH'`: which boundary data DUMPFLUX writes (default `'FLUX'`)"],
    ["USEFLUX", "GRID", "Run one sector only, with the boundary from a FLUX file; optional file name"],
    ["FLUXREG", "GRID", "Which FLUXNUM region the USEFLUX run covers"],
    ["FLUXALL", "SUMMARY", "In the parent: report the summary vectors a reduced run needs"],
    ["FPRFLUX", "SUMMARY", "In a reduced run: field pressure over the whole parent model (automatic)"],
], LEFT, TOP + Inches(0.1), FULL_W, [1.3, 1.2, 8.5], size=17, row_h=0.52)
caption(s, "One region: `CASE.FLUX`.  Several regions: `CASE.FLUX0001`, `CASE.FLUX0002`, …  "
        "FLUXNUM without DUMPFLUX is just a region array.",
        LEFT, TOP + Inches(4.6), FULL_W, size=16, align=PP_ALIGN.LEFT)

s = new_slide("Producing the FLUX file: DUMPFLUX or make_flux")
cw, x1, x2 = two_columns()
card(s, x1, TOP, cw, Inches(5.3), "DUMPFLUX, during the parent run", [
    "Keyword in the parent's GRID section",
    "Boundary written at **every time step** of the parent",
    "One file per FLUXNUM region",
    "Most accurate boundary",
    "Has to be decided **before** the parent run",
], BLUE_CARD, size=21)
card(s, x2, TOP, cw, Inches(5.3), "make_flux, after the parent run", [
    "Command-line tool, applied to an ordinary run",
    "Reads the deck, EGRID, INIT, restart and summary files",
    "Sector from a box or cell list on the command line, or a FLUXNUM file",
    "**No rerun** of the full model",
    "Boundary only at the parent's **restart steps**: ask for a restart at "
    "every report step for good accuracy",
], ORANGE_CARD, size=21)
notes(s, """Both produce the same file format, and a reduced run cannot tell
them apart. make_flux is the tool for existing runs; how good its boundary is
depends on how often the parent wrote restarts, as example 2 shows.""")

s = new_slide("Running the sector: parent deck or cut-out deck")
cw, x1, x2 = two_columns()
card(s, x1, TOP, cw, Inches(5.3), "Reduced run on the parent deck", [
    "Same deck: DUMPFLUX replaced by **USEFLUX** and **FLUXREG**",
    "Everything outside the region is made inactive",
    "**No schedule changes**: wells outside the sector simply have no connections",
    "The simplest way to get started",
], GREEN_CARD, size=21)
card(s, x2, TOP, cw, Inches(5.3), "Reduced run on a cut-out sector deck", [
    "A deck of its own, with a grid that holds only the sector",
    "The grid is the sector's bounding box, with the parent's cells",
    "FLUXNUM is 1 everywhere; USEFLUX + FLUXREG 1",
    "Keeps only its own wells: the parent's other wells are **added "
    "automatically** from the FLUX file",
], PINK_CARD, size=21)
notes(s, """Either way the reduced run simulates the same cells. The cut-out
deck is smaller to read and process and is the usual starting point for a
detailed sector study; the parent deck needs no preparation at all.""")

s = new_slide("Pressure boundary or flux boundary")
cw, x1, x2 = two_columns()
card(s, x1, TOP, cw, Inches(4.55), "Pressure  ·  FLUXTYPE 'PRESSURE'", [
    "The file holds the state beyond every boundary face: pressure, "
    "saturations, Rs, Rv, kr and Pc",
    "Flow across the face is computed with the parent's transmissibility",
    "**Responds to changes** inside the sector: a new well near the "
    "boundary draws in more fluid",
    "The mode for studying changes",
], BLUE_CARD, size=20)
card(s, x2, TOP, cw, Inches(4.55), "Flux  ·  FLUXTYPE 'FLUX'", [
    "The file holds the mass of oil, gas and water crossing every face",
    "The rates are imposed as they are",
    "Reproduces the parent most closely; smaller files",
    "**Does not respond** to changes inside the sector",
], ORANGE_CARD, size=20)
bullets(s, [
    "FLUXTYPE is set in the **parent**; the reduced run follows whatever the file holds",
    "`'BOTH'` writes both kinds of data; a reduced run then imposes the rates",
], y=TOP + Inches(4.7), h=Inches(0.9), size=19, gap=2)

s = new_slide("Boundary details handled for you")
bullets(s, [
    "Boundary records at **every time step** of the parent, interpolated in time by the reduced run",
    "Rates are **time averages** of the converged solution: the boundary moves "
    "exactly the mass the parent moved",
    "Exterior pressure carried **hydrostatically** from the cell centre to the face",
    "Incoming fluid uses the PVT region, kr and Pc of the **cell it comes from**",
    "Faces across **faults with throw, pinch-outs and other NNCs** are part of the boundary; "
    "faces towards inactive cells stay closed",
    "**Threshold pressures** (THPRES) are honoured, with the values the parent arrived at",
    "Serial and **parallel** runs write the same FLUX file and give the same results",
], size=22, gap=10)
notes(s, """These details are why the reduced run agrees with the parent to a
few tenths of a percent rather than a few percent. None of them needs any
input from the user.""")

s = new_slide("The rest of the field: wells, groups and UDQs")
cw, x1, x2 = two_columns()
bullets(s, [
    "#The FLUX file also carries",
    "**Summary vectors** of the parent: all well rates and cumulatives, and "
    "everything the UDQs and ACTIONX conditions refer to",
    "**Wells and groups** of the parent: group, type, status, efficiency "
    "factor and the group tree, per report step",
], x=x1, w=cw, size=21, gap=8)
bullets(s, [
    "#So that in the reduced run",
    "Outside wells take part in **group control** with their parent rates; "
    "the wells inside get the share of group targets they had",
    "**FOPR, FOPT** and group vectors are whole-field values, comparable with the parent",
    "**UDQs, UDAs and ACTIONX** see the parent's values for whatever the sector cannot compute",
    "A cut-out deck gets the parent's missing wells **added automatically**",
], x=x2, w=cw, size=21, gap=8)
notes(s, """Group controls, UDQs and ACTIONX are normally written for the
whole field. The reduced run uses the parent's values for the outside of the
sector, so that the wells inside behave as they did in the parent.""")

s = new_slide("FLUXALL and FPRFLUX")
cw, x1, x2 = two_columns()
card(s, x1, TOP, cw, Inches(5.0), "FLUXALL  ·  parent SUMMARY", [
    "Reports exactly the summary vectors a reduced run needs",
    "Which ones depends on the wells, and on what the UDQs and ACTIONX "
    "conditions refer to: no need to list them by hand",
    "**Required for make_flux**, which can only embed what the parent wrote",
    "Recommended for DUMPFLUX",
], BLUE_CARD, size=20)
card(s, x2, TOP, cw, Inches(5.0), "FPRFLUX  ·  reduced run", [
    "Added automatically to every USEFLUX run",
    "Field pressure averaged over the **whole parent model**: the parent's "
    "sums outside the sector plus the sector's own cells",
    "FPR of a reduced run covers the sector only: compare **FPRFLUX** with the parent's FPR",
    "The same whole-model averages convert rates of wells on RESV control",
], GREEN_CARD, size=20)

# ------------------------------------------------------------------ Examples

s = dark_slide("Worked examples", "Sector modelling on the Drogon model", image="full_pressure.png")

s = new_slide("The Drogon test case")
bullets(s, [
    "Synthetic three-phase black-oil model with dissolved gas and vaporised oil",
    "Faults, 7 equilibration regions with threshold pressures",
    "46 × 73 × 32 grid, **72 339** active cells; history followed by prediction",
    "Sector: box I 10–41, J 37–64, K 1–32: **17 499** cells, a quarter of the model",
    "Inside: producers **A3, A4**, injector **A6**",
    "Outside: producers A1, A2, injector A5",
], w=Inches(6.6), size=19)
table(s, [
    ["Run", "Cells", "Wall time"],
    ["Parent `DUMP-SECTOR`", "72 339", "105 s"],
    ["Reduced on parent deck", "17 499", "27 s"],
    ["Reduced, cut-out deck", "17 499", "26 s"],
], LEFT, TOP + Inches(4.0), Inches(6.4), [3.2, 1.2, 1.4], size=15, row_h=0.36, align="lrr")
picture(s, "full_fluxnum.png", Inches(7.4), TOP + Inches(0.2), Inches(5.3), Inches(3.9))
caption(s, "FLUXNUM: sector (red) and the rest (blue)", Inches(7.4), TOP + Inches(4.2), Inches(5.3))
notes(s, """Wall times are for four threads, runs made one after the other on
the same machine. Comparisons in the following slides give the largest
deviation over the run, relative to the largest value of the parent's
vector.""")

s = new_slide("Example 1  ·  Parent run with DUMPFLUX")
cw, x1, x2 = two_columns()
caption(s, "GRID section of `DUMP-SECTOR.DATA`", x1, TOP - Inches(0.05), cw, size=15, align=PP_ALIGN.LEFT)
code(s, """DUMPFLUX

FLUXTYPE
  'PRESSURE' /

EQUALS
  FLUXNUM 0 /
  FLUXNUM 1 10 41 37 64 1 32 /
/""", x1, TOP + Inches(0.35), cw, Inches(2.9), size=15)
caption(s, "SUMMARY section", x1, TOP + Inches(3.4), cw, size=15, align=PP_ALIGN.LEFT)
code(s, """SUMMARY
ALL
PERFORMA
FLUXALL""", x1, TOP + Inches(3.8), cw, Inches(1.5), size=15)
caption(s, "Run it as usual", x2, TOP - Inches(0.05), cw, size=15, align=PP_ALIGN.LEFT)
code(s, "flow DUMP-SECTOR.DATA", x2, TOP + Inches(0.35), cw, Inches(0.5), size=15)
caption(s, "The print file reports", x2, TOP + Inches(1.05), cw, size=15, align=PP_ALIGN.LEFT)
code(s, """DUMPFLUX bootstrap: initialized 1 region
  dumper(s) from FLUXNUM
DUMPFLUX will write sector boundary data
  every time step, at most every 1 day(s)
DUMPFLUX will embed 176 summary vector(s),
  sampled at most every 1 day(s)""", x2, TOP + Inches(1.45), cw, Inches(2.0), size=13)
bullets(s, [
    "Writes `DUMP-SECTOR.FLUX`, **34 MB**",
    "2427 boundary faces, 157 records, 176 summary vectors",
], x=x2, y=TOP + Inches(3.65), w=cw, h=Inches(1.6), size=18)

s = new_slide("Example 1  ·  Reduced run with USEFLUX")
cw, x1, x2 = two_columns()
caption(s, "GRID section of `USE-SECTOR.DATA`", x1, TOP - Inches(0.05), cw, size=15, align=PP_ALIGN.LEFT)
code(s, """USEFLUX
  'DUMP-SECTOR' /

FLUXTYPE
  'PRESSURE' /

FLUXREG
  1 /

EQUALS
  FLUXNUM 0 /
  FLUXNUM 1 10 41 37 64 1 32 /
/""", x1, TOP + Inches(0.35), cw, Inches(3.75), size=15)
code(s, "flow USE-SECTOR.DATA", x1, TOP + Inches(4.3), cw, Inches(0.5), size=15)
bullets(s, [
    "USEFLUX names the FLUX file, without `.FLUX`",
    "FLUXREG selects the region",
    "FLUXTYPE is ignored: **the file decides** the boundary type",
    "FLUXALL can be left out",
    "**No changes to SCHEDULE**",
], x=x2, w=cw, h=Inches(2.7), size=18)
caption(s, "The print file reports", x2, TOP + Inches(2.85), cw, size=15, align=PP_ALIGN.LEFT)
code(s, """USEFLUX: 157 boundary record(s) read from
  'DUMP-SECTOR.FLUX' (per time step),
  interpolated in time between records
USEFLUX: summary vectors loaded from flux
  file 'DUMP-SECTOR.FLUX' (176 keys, ...)
USEFLUX: adopted the producing run's
  threshold pressures over 7 equilibration
  region(s)""", x2, TOP + Inches(3.25), cw, Inches(2.3), size=12)

s = new_slide("Example 1  ·  Pressure at the end of the run")
cw, x1, x2 = two_columns(Inches(0.3))
picture(s, "full_pressure.png", x1, TOP + Inches(0.2), cw, Inches(4.6))
picture(s, "usesector_pressure.png", x2, TOP + Inches(0.2), cw, Inches(4.6))
caption(s, "Parent: the whole field", x1, TOP + Inches(4.9), cw, size=16)
caption(s, "Reduced run on the parent deck: only the sector is active", x2, TOP + Inches(4.9), cw, size=16)

s = new_slide("Example 1  ·  Flux boundaries")
caption(s, "Only FLUXTYPE in the parent changes, and the reduced deck names the new file",
        LEFT, TOP - Inches(0.05), FULL_W, size=17, align=PP_ALIGN.LEFT)
code(s, """-- parent, DUMP-SECTOR-FLUX.DATA       -- reduced, USE-SECTOR-FLUX.DATA
DUMPFLUX                               USEFLUX
                                         'DUMP-SECTOR-FLUX' /
FLUXTYPE                               FLUXREG
  'FLUX' /                               1 /""", LEFT, TOP + Inches(0.45), Inches(9.2), Inches(1.75), size=15)
bullets(s, [
    "A flux-mode file only holds component mass rates per face: **9.7 MB** instead of 34 MB",
    "The reduced run imposes the parent's rates: the closest match, but no response to changes",
    "Results of both modes on the next slides",
], y=TOP + Inches(2.5), h=Inches(2.5), size=19)

s = new_slide("Example 1  ·  Results")
table(s, [
    ["Vector", "Parent (end)", "Pressure", "Flux"],
    ["FOPR", "1883.5", "0.31 %", "0.31 %"],
    ["FOPT", "1.341e7", "0.02 %", "0.00 %"],
    ["FPR", "252.73", "4.18 %", "4.14 %"],
    ["FPRFLUX (parent FPR)", "252.73", "0.03 %", "0.02 %"],
    ["WOPR:A3", "406.59", "0.17 %", "0.07 %"],
    ["WOPT:A3", "2.361e6", "0.04 %", "0.02 %"],
    ["WBHP:A3", "170.53", "0.17 %", "0.06 %"],
    ["WTHP:A3", "0", "0.15 %", "0.07 %"],
    ["WOPR:A4", "605.79", "0.18 %", "0.08 %"],
    ["WOPT:A4", "4.608e6", "0.04 %", "0.00 %"],
    ["WBHP:A4", "223.38", "0.02 %", "0.05 %"],
    ["WTHP:A4", "0", "0.12 %", "0.10 %"],
    ["WWIR:A6", "6500", "0.00 %", "0.00 %"],
    ["WWIT:A6", "1.536e7", "0.00 %", "0.00 %"],
    ["WBHP:A6", "336.44", "0.08 %", "0.09 %"],
], LEFT, TOP - Inches(0.05), Inches(7.0), [2.9, 1.6, 1.25, 1.25], size=13, row_h=0.345, align="lrrr")
bullets(s, [
    "Largest deviation over the run, relative to the largest value of the parent",
    "Rates, cumulatives and pressures of the in-sector wells within **0.2 %**",
    "FPR covers the sector only, hence 4 %; **FPRFLUX** matches the parent's FPR",
    "WWIR of A6 is set by its control; its BHP is not",
], x=Inches(8.0), w=Inches(4.75), size=18)

s = new_slide("Example 1  ·  Field oil production")
picture(s, "ex1_field_oil.png", LEFT, TOP, FULL_W, Inches(3.5))
bullets(s, [
    "FOPR and FOPT are **field totals**: they include A1 and A2, outside the sector, at their parent rates",
    "The 0.3 % deviation in FOPR is confined to a few steps around day 865, "
    "when the producers are briefly shut in",
    "Elsewhere the reduced runs agree to within a few thousandths of a percent",
], y=TOP + Inches(3.6), h=Inches(2.0), size=18, gap=3)

s = new_slide("Example 1  ·  FPR and FPRFLUX")
picture(s, "ex1_field_pressure.png", LEFT, TOP, FULL_W, Inches(3.5))
bullets(s, [
    "**FPR** of a reduced run averages over the cells it simulates, the sector "
    "alone: about 12 bar off the parent's FPR at the end",
    "**FPRFLUX** adds the parent's pore-volume weighted sums outside the sector: "
    "the whole-model average, within **0.03 %**",
    "Compare FPRFLUX with the parent's FPR",
], y=TOP + Inches(3.6), h=Inches(2.0), size=18, gap=3)

s = new_slide("Example 1  ·  Producer A3")
picture(s, "ex1_well_A3.png", LEFT, TOP - Inches(0.1), FULL_W, Inches(5.7))

s = new_slide("Example 1  ·  Producer A4")
picture(s, "ex1_well_A4.png", LEFT, TOP - Inches(0.1), FULL_W, Inches(5.7))

s = new_slide("Example 1  ·  Injector A6")
picture(s, "ex1_well_A6.png", LEFT, TOP - Inches(0.1), FULL_W, Inches(5.7))
notes(s, """The injection rate is set by the well's control, which is why it is
identical; the bottom-hole pressure is computed and agrees to 0.1 %.""")

s = new_slide("Example 2  ·  make_flux and a cut-out sector deck")
bullets(s, [
    "The FLUX file is built **afterwards** from the parent's output",
    "The reduced run uses a deck with a grid holding **only the sector**",
], y=TOP - Inches(0.1), h=Inches(1.1), size=19, gap=3)
caption(s, "Build the FLUX file from the output of `DUMP-SECTOR` (its own FLUX file is ignored)",
        LEFT, TOP + Inches(1.05), FULL_W, size=16, align=PP_ALIGN.LEFT)
code(s, """make_flux --parent=DUMP-SECTOR --mapping-inline='box 10 37 1 41 64 32' \\
          --output=TEST.FLUX""", LEFT, TOP + Inches(1.45), FULL_W, Inches(0.8), size=15)
caption(s, "make_flux reports", LEFT, TOP + Inches(2.4), FULL_W, size=16, align=PP_ALIGN.LEFT)
code(s, """Embedding 176 of the parent's 872 summary vectors, drawn from 16 keyword(s) a reduced run can use
Note: the parent wrote 6 restart step(s) over 121 report step(s), so the boundary is
      described at that cadence. The summary is embedded at its own.
      Rerun the parent with RPTRST BASIC=2 for a record at every report step.
Wrote pressure-mode FLUX file 'TEST.FLUX' for region 1 with 2427 boundary faces,
      6 report steps and 156 summary samples""", LEFT, TOP + Inches(2.8), FULL_W, Inches(1.9), size=13)
caption(s, "Box: `box i1 j1 k1 i2 j2 k2`, one-based and inclusive.  Pressure mode is the default.",
        LEFT, TOP + Inches(4.85), FULL_W, size=16, align=PP_ALIGN.LEFT)
notes(s, """Any run can be the parent of make_flux. Here the output of the
DUMPFLUX run of example 1 is reused, so that the results can be compared
directly. Note the hint about restart output: it matters, as the following
slides show.""")

s = new_slide("Example 2  ·  The cut-out sector deck")
cw = Inches(5.6)
caption(s, "Added to the GRID section of `SECTOR-UF.DATA`", LEFT, TOP - Inches(0.05), cw,
        size=15, align=PP_ALIGN.LEFT)
code(s, """USEFLUX
  '../../TEST' /

FLUXTYPE
  'PRESSURE' /

FLUXREG
  1 /

EQUALS
  FLUXNUM 1 /
/""", LEFT, TOP + Inches(0.35), cw, Inches(3.5), size=15)
bullets(s, [
    "32 × 28 × 32 grid cut out of the parent",
    "FLUXNUM 1 everywhere: the whole grid is the sector",
    "File name relative to the deck",
], y=TOP + Inches(3.95), w=cw, h=Inches(1.6), size=17, gap=2)
picture(s, "cutout_pressure.png", LEFT + cw + Inches(0.3), TOP - Inches(0.1),
        FULL_W - cw - Inches(0.3), Inches(3.9))
code(s, """USEFLUX: 5 well(s) of the parent run are not in
  this deck and have been added without connections,
  their rates taken from '../../TEST.FLUX':
  R_A2 R_A5 A1 A2 A5""", LEFT + cw + Inches(0.3), TOP + Inches(3.95),
     FULL_W - cw - Inches(0.3), Inches(1.45), size=13)
notes(s, """The deck only has the wells inside the sector. The other wells are
added from the FLUX file, without connections, and their rates are the
parent's. In the picture they are the labels at the corner of the grid: the
added wells are placed at cell (1, 1).""")

s = new_slide("Example 2  ·  Restart output matters")
bullets(s, [
    "make_flux writes a boundary record **per restart** of the parent",
    "Drogon's base case writes 6 restarts in 7 years",
    "Rerun the parent with a restart at every report step:",
], x=LEFT, w=Inches(4.4), h=Inches(2.6), size=17, gap=4)
code(s, """RPTRST
 BASIC=2 FLOWS FLORES /""", LEFT, TOP + Inches(2.55), Inches(4.4), Inches(0.75), size=14)
bullets(s, [
    "91 records: within **0.5 %** in rates and **1 %** in pressures",
    "The cut-out grid costs nothing: on the DUMPFLUX file it matches example 1",
], x=LEFT, y=TOP + Inches(3.45), w=Inches(4.4), h=Inches(2.1), size=17, gap=4)
table(s, [
    ["Vector", "Parent", "6 restarts", "91 restarts", "DUMPFLUX"],
    ["FOPR", "1883.5", "1.01 %", "0.31 %", "0.31 %"],
    ["FOPT", "1.341e7", "0.56 %", "0.14 %", "0.02 %"],
    ["FPRFLUX", "252.73", "2.62 %", "0.21 %", "0.03 %"],
    ["WOPR:A3", "406.59", "2.89 %", "0.26 %", "0.17 %"],
    ["WOPT:A3", "2.361e6", "1.50 %", "0.46 %", "0.05 %"],
    ["WBHP:A3", "170.53", "7.08 %", "0.71 %", "0.17 %"],
    ["WTHP:A3", "0", "9.55 %", "0.10 %", "0.15 %"],
    ["WOPR:A4", "605.79", "1.98 %", "0.23 %", "0.18 %"],
    ["WOPT:A4", "4.608e6", "0.87 %", "0.17 %", "0.04 %"],
    ["WBHP:A4", "223.38", "4.58 %", "0.43 %", "0.02 %"],
    ["WTHP:A4", "0", "7.18 %", "0.14 %", "0.12 %"],
    ["WWIR:A6", "6500", "0.00 %", "0.00 %", "0.00 %"],
    ["WBHP:A6", "336.44", "2.81 %", "0.20 %", "0.08 %"],
], Inches(5.3), TOP - Inches(0.05), Inches(7.43), [1.7, 1.3, 1.3, 1.4, 1.3], size=13,
    row_h=0.37, align="lrrrr")
notes(s, """The three columns are the same cut-out deck on three FLUX files:
make_flux on the base case with six restarts, make_flux on a rerun with a
restart at every report step, and the DUMPFLUX file of example 1, which has a
record at every time step. FPR is left out here; as in example 1 it covers
the sector only, and FPRFLUX is the one to compare.""")

s = new_slide("Example 2  ·  Field results of the cut-out sector")
picture(s, "ex2_field_oil.png", LEFT, TOP - Inches(0.1), FULL_W, Inches(2.85))
picture(s, "ex2_field_pressure.png", LEFT, TOP + Inches(2.8), FULL_W, Inches(2.85))
notes(s, """With only six records the field pressure follows the parent's
decline only between the restart steps. With a record at every report step,
the curves lie on top of the parent's.""")

s = new_slide("Example 2  ·  Producer A3 in the cut-out sector")
picture(s, "ex2_well_A3.png", LEFT, TOP - Inches(0.1), FULL_W, Inches(5.7))
notes(s, """The bottom-hole and tubing-head pressures are the most sensitive
to the boundary. With six restarts the error reaches seven to ten percent;
with a restart at every report step it is below one percent.""")

s = new_slide("Example 2  ·  Producer A4 in the cut-out sector")
picture(s, "ex2_well_A4.png", LEFT, TOP - Inches(0.1), FULL_W, Inches(5.7))

s = new_slide("Example 2  ·  Injector A6 in the cut-out sector")
picture(s, "ex2_well_A6.png", LEFT, TOP - Inches(0.1), FULL_W, Inches(5.7))

s = new_slide("Example 3  ·  UDQs over wells outside the sector")
cw, x1, x2 = two_columns()
caption(s, "Added to parent and cut-out deck, before the prediction", x1, TOP - Inches(0.05), cw,
        size=15, align=PP_ALIGN.LEFT)
code(s, """UDQ
  DEFINE FUOUTOPR  WOPR 'A1' + WOPR 'A2' /
  UNITS  FUOUTOPR  SM3/DAY /
  DEFINE FUSECOPR  WOPR 'A3' + WOPR 'A4' /
  UNITS  FUSECOPR  SM3/DAY /
/""", x1, TOP + Inches(0.35), cw, Inches(2.2), size=14)
bullets(s, [
    "**FUOUTOPR** only uses wells outside the sector, which the cut-out deck does not even define",
    "FLUXALL makes the parent report WOPR of every well the UDQs use",
    "FLUX file built with make_flux as before",
], x=x1, y=TOP + Inches(2.7), w=cw, h=Inches(2.8), size=17, gap=4)
table(s, [
    ["Vector", "Parent (end)", "Cut-out sector"],
    ["FOPR", "1883.5", "0.31 %"],
    ["FOPT", "1.341e7", "0.14 %"],
    ["FPRFLUX", "252.73", "0.21 %"],
    ["FUOUTOPR", "871.1", "0.00 %"],
    ["FUSECOPR", "1012.4", "0.47 %"],
], x2, TOP, cw, [2.0, 1.6, 1.8], size=15, row_h=0.42, align="lrr")
bullets(s, [
    "FUOUTOPR is **exact**: A1 and A2 have the parent's rates",
    "FUSECOPR deviates as much as A3 and A4 do",
], x=x2, y=TOP + Inches(2.75), w=cw, h=Inches(2.5), size=17, gap=4)

s = new_slide("Example 3  ·  The UDQs in parent and sector")
picture(s, "ex3_extra.png", LEFT, TOP, FULL_W, Inches(3.5))
bullets(s, [
    "A1 and A2 are added to the sector's schedule from the FLUX file, after "
    "the deck is read: the UDQ check for undefined wells becomes a warning",
    "UDQs used as group targets and ACTIONX conditions work the same way: the "
    "reduced run sees the parent's values for whatever it cannot compute",
], y=TOP + Inches(3.6), h=Inches(2.0), size=18, gap=3)

# -------------------------------------------------------------- Wrapping up

s = new_slide("Practical checklist")
cw, x1, x2 = two_columns()
bullets(s, [
    "#Parent",
    "Define FLUXNUM; keep every well within one region",
    "Add **FLUXALL** to SUMMARY",
    "DUMPFLUX with **FLUXTYPE 'PRESSURE'** for studies of changes",
    "For make_flux: a **restart at every report step**, with FLOWS for flux mode",
], x=x1, w=cw, size=21, gap=8)
bullets(s, [
    "#Reduced run",
    "USEFLUX naming the file, FLUXREG the region",
    "Check the USEFLUX lines in the print file: records, summary keys, added wells",
    "Compare **FPRFLUX**, not FPR, with the parent's FPR",
    "Keep changes away from the boundary: the outside does not respond",
    "`inspect_flux` shows what a FLUX file holds for a cell or a summary vector",
], x=x2, w=cw, size=21, gap=8)

s = new_slide("Limitations")
bullets(s, [
    "**Black-oil only**; no solvent, polymer, brine or tracers across the boundary",
    "**One sector per reduced run**, and each well within a single FLUXNUM region",
    "**No feedback to the outside**: the boundary and the outside wells follow the parent",
    "**Same grid geometry**: the parent grid or the sector's bounding box, no refinement or coarsening",
    "**make_flux** is limited by the parent's output: restart cadence, no hysteresis, "
    "FLOWS needed for flux mode",
    "Restarted runs and LGRs are **untested**",
    "**File size** grows with faces × time steps: 34 MB for Drogon in pressure mode",
], size=22, gap=10)
notes(s, """Details are in the limitations chapter of the guide. The most
important one in practice is that the outside does not respond to changes in
the sector: the larger the change, the further from the boundary it should
be.""")

s = new_slide("Summary")
bullets(s, [
    "Run a sector of a field model with boundary data from a parent run",
    "DUMPFLUX during the parent run, or **make_flux** afterwards",
    "On the parent deck, or on a **cut-out sector deck**",
    "Pressure boundaries respond to changes, flux boundaries reproduce the parent",
    "Wells, groups, UDQs and ACTIONX outside the sector follow the parent",
    "Drogon: within a few tenths of a percent, about **4 × faster**",
    "#Further reading",
    "The guide: `opm-simulators/doc/sector`, built with `make sector_docs` "
    "as HTML and PDF",
    "Reproduce the examples: `doc/sector/tools/run-examples.sh`",
], size=21, gap=8)

# -------------------------------------------------------- Technical, optional

s = dark_slide("Technical details", "Optional slides\nImplementation, file format and coupling to the parent")

s = new_slide("Implementation overview", optional=True)
table(s, [
    ["Component", "Where", "Role"],
    ["`FluxFile`", "opm-common", "Reads and writes the FLUX file format"],
    ["`IOConfig`, `SummaryConfig`", "opm-common", "Flux keywords; FLUXALL expansion, FPRFLUX"],
    ["`Schedule::insertKeywords`", "opm-common", "Adds the parent's wells to a reduced run's schedule"],
    ["`FluxRegions`", "`flow/flux/`", "Cells and boundary faces of each FLUXNUM region"],
    ["`FluxDumper`", "`flow/flux/`", "Collects and appends one region's boundary data"],
    ["`FluxBoundary`", "`flow/flux/`", "Maps a FLUX file onto the reduced run's grid"],
    ["`FluxActivation`", "`flow/flux/`", "Narrows ACTNUM to the selected region"],
    ["`ParentSummary`", "`flow/flux/`", "Time series of the parent's summary vectors"],
    ["`FluxSummaryKeys`, `FluxParentWells`", "`flow/flux/`", "What is embedded; the parent's wells"],
    ["`EclWriter`", "`flow/`", "The DUMPFLUX producer; parent values in the summary"],
    ["`FlowProblem`, `FlowProblemBlackoil`", "`flow/`", "USEFLUX boundary conditions, parent coupling"],
    ["`make_flux`, `inspect_flux`", "`examples/`", "Offline producer and inspection tool"],
], LEFT, TOP - Inches(0.05), FULL_W, [3.9, 1.7, 6.3], size=13, row_h=0.37)
caption(s, "`flow/` is `opm/simulators/flow` in opm-simulators", LEFT, TOP + Inches(5.3), FULL_W,
        size=13, align=PP_ALIGN.LEFT)

s = new_slide("How DUMPFLUX records the boundary", optional=True)
cw, x1, x2 = two_columns()
bullets(s, [
    "#Boundary faces",
    "Cartesian faces between a region cell and an active cell outside",
    "NNCs from the deck, pinch-outs and faults with throw",
    "Each well must lie within one region, or the run stops",
    "#Rates (flux mode)",
    "Converged FLORES per face, turned into component masses with the "
    "upstream cell's b, Rs, Rv and reference densities",
    "Checked against FLOWS; averaged over the record's interval",
], x=x1, w=cw, size=18, gap=5)
bullets(s, [
    "#Exterior state (pressure mode)",
    "Pressure, Sw, Sg, Rs, Rv, T at the end of the interval",
    "kr and Pc of the exterior cell from the converged state: "
    "saturation region, scaling and hysteresis included",
    "Static: transmissibility, PVT and equilibration region, depth",
    "#Also recorded",
    "Pore-volume weighted sums outside the region, threshold pressures, "
    "parent wells per report step, summary samples",
    "Appended as formed; reduced to the I/O rank in parallel",
], x=x2, w=cw, size=18, gap=5)

s = new_slide("The FLUX file", optional=True)
bullets(s, [
    "Eclipse-style keyword file, binary or formatted (`.FFLUX`); readable with "
    "`opm.io.ecl.EclFile`, convertible with `convertECL`; SI units, times in seconds",
], y=TOP - Inches(0.1), h=Inches(0.9), size=17)
cw = (FULL_W - Inches(0.6)) // 3
blocks = [
    ("Static section", ["`FLUXHEAD`: version 7, grid, box, mode, phases",
                        "`LOC2GLOB`: cells of the region",
                        "`FACECELL`, `FACEDIR`, `FACEGLNB`: faces",
                        "`FLUXTRAN`, `FLXPVTN`, `FLXEXDP`, `FLXEQLN`",
                        "`FLXTHPR`: threshold pressures",
                        "`FLXWELLS`, `FLXWGRP`, …: parent wells"], BLUE_CARD),
    ("Record blocks", ["`FLXSEQ`: block number, records",
                       "`FLXTIME`, `FLXDT`: interval",
                       "`FLXMASS`: mass rates (flux)",
                       "`FLXPRES`, `FLXSATW`, `FLXSATG`, `FLXRS`, `FLXRV` (pressure)",
                       "`FLXKR`, `FLXPC`",
                       "`FLXRCON`: sums outside"], ORANGE_CARD),
    ("Summary blocks", ["`SMRYSEQ`: block number, samples",
                        "`SMRYTIME`: end of each interval",
                        "`SMRYVALS`: values, sample-major",
                        "Keys in `FLUXNAMS`",
                        "Rates are interval averages"], GREEN_CARD),
]
for i, (head, items, fill) in enumerate(blocks):
    card(s, LEFT + i * (cw + Inches(0.3)), TOP + Inches(0.85), cw, Inches(4.25), head, items, fill, size=16)
caption(s, "FACEDIR: 1 I+, 2 I−, 4 J+, 8 J−, 16 K+, 32 K−, 0 NNC.   Mode: 1 flux, 2 pressure, 3 both.   "
        "Phase mask: 1 oil, 2 water, 4 gas.", LEFT, TOP + Inches(5.2), FULL_W, size=13, align=PP_ALIGN.LEFT)

s = new_slide("Boundary conditions in the reduced run", optional=True)
_, tf = textbox(s, LEFT, TOP - Inches(0.05), FULL_W, Inches(0.7))
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
add_runs(p, "*q*_{α} = *T* · *λ*_{α}^{up} · ( *p*_{α}^{ext} − *p*_{α}^{int} + *ρ̄*_{α} *g* (*z*^{int} − *z*^{face}) )",
         24, NAVY)
cw, x1, x2 = two_columns()
bullets(s, [
    "#Pressure boundary",
    "*T*: the parent's transmissibility across the face",
    "Exterior fluid state from the record, with the exterior cell's PVT region",
    "Pressure carried hydrostatically from the exterior cell centre to the face",
    "Incoming fluid uses the exterior's recorded kr",
    "Threshold pressure between the equilibration regions, from the parent",
], x=x1, y=TOP + Inches(0.8), w=cw, h=Inches(4.7), size=19, gap=5)
bullets(s, [
    "#Flux boundary",
    "Recorded mass rate per component and face, imposed as is",
    "#NNC faces",
    "No grid face in the reduced run: the flux becomes a source term in the interior cell",
    "#In time",
    "Record covering the start of each step; states interpolated linearly to the step's end",
    "Rates are interval averages, not interpolated",
], x=x2, y=TOP + Inches(0.8), w=cw, h=Inches(4.7), size=19, gap=5)

s = new_slide("Wells outside the sector", optional=True)
cw, x1, x2 = two_columns()
bullets(s, [
    "#Wells with no connections in the reduced grid",
    "Never solved for",
    "**Summary**: rates over the step and BHP, THP at its end taken from the "
    "parent summary, as if simulated; field and group totals formed as usual",
    "**Group control**: parent surface and reservoir rates as fixed rates; "
    "they count towards groups and come off targets, but get no share",
    "Parent summary: rates piecewise constant, other vectors interpolated linearly",
], x=x1, w=cw, size=19, gap=6)
bullets(s, [
    "#Wells missing from a cut-out deck",
    "Generated where something changes, per report step:",
    (1, "GRUPTREE, GEFAC for missing groups"),
    (1, "WELSPECS in the parent group, head at (1, 1)"),
    (1, "WCONPROD or WCONINJE, parent status, zero target"),
    (1, "WEFAC with the parent's efficiency factor"),
    "Inserted with `Schedule::insertKeywords()`, one rebuild from the first affected step",
    "UDQ and WELLDIMS checks become warnings while reading",
], x=x2, w=cw, size=19, gap=6)

s = new_slide("Summary vectors and UDQs in the reduced run", optional=True)
bullets(s, [
    "**Start of every step**: the parent's values of all embedded "
    "non-cumulative vectors are written into the summary state, for UDAs, "
    "group controls and ACTIONX",
    "Cumulatives are not seeded; they accumulate from the rates",
    "**After the summary evaluation, before the UDQs**: the parent's values are put back for",
    (1, "well, connection and segment vectors of wells outside the sector"),
    (1, "field and group vectors that are not sums of well rates"),
    (1, "region vectors, and block vectors of inactive cells"),
    "UDQs over regions outside the sector are accepted; region sets are extended "
    "with the regions the parent reported",
    "A UDQ group target keeps its start-of-step value, the same on every rank",
    "FPRFLUX and RESV rate conversion add the parent's outside sums to the sector's own",
], size=20, gap=7)

s = new_slide("make_flux in more detail", optional=True)
bullets(s, [
    "**Region**: `--fluxnum` grdecl file with `--regions`, or `--mapping` / "
    "`--mapping-inline` with boxes, cells or index ranges",
    "**Records**: one per restart step; pressure, saturations, Rs, Rv from the restart",
    "**kr and Pc**: evaluated from the parent's saturation functions and scaled end "
    "points; drainage curves, since hysteresis is not in the restart",
    "**Flux mode**: FLOWS arrays × reference density; needs a mapping, not `--fluxnum`",
    "**Threshold pressures** from THRESHPR in the restart; a warning if THPRES is "
    "used but THRESHPR missing (e.g. `NORST=1`)",
    "**Summary**: as DUMPFLUX, from the summary file; `--summary-min-interval` "
    "thins it, `--smry-define` supplies a missing vector",
], size=21, gap=10)

s = new_slide("Run-time parameters and inspection", optional=True)
table(s, [
    ["Parameter", "Default", "Effect"],
    ["`--flux-boundary-min-interval-between-samples`", "1 day", "Minimum time between boundary records; 0 every step"],
    ["`--flux-boundary-report-steps-only`", "false", "Boundary records at report steps only"],
    ["`--flux-summary-min-interval-between-samples`", "1 day", "Minimum time between summary samples"],
    ["`--flux-boundary-interpolate`", "true", "Interpolate the boundary state in time"],
], LEFT, TOP - Inches(0.05), FULL_W, [5.2, 1.0, 5.0], size=13, row_h=0.42)
caption(s, "`inspect_flux` prints the boundary of one cell as a time series, or embedded summary vectors",
        LEFT, TOP + Inches(2.3), FULL_W, size=15, align=PP_ALIGN.LEFT)
code(s, """$ inspect_flux DUMP-SECTOR.FLUX 10 50 5 --start-date=2018-01-01
  parent case      DUMP-SECTOR
  records          157
  I-     exterior cell     15694       (9,50,5)  transmissibility 4.664e-13 m3.Pa.s  pvt region 1
#       DATE      I-:PRES      I-:SWAT      I-:SGAS        I-:RS        I-:RV  ...
  2018-01-01    3.015e+07       0.6966            0        140.8    0.0002351  ...""",
     LEFT, TOP + Inches(2.75), FULL_W, Inches(1.85), size=12)
caption(s, "Tests: `fluxTests.cmake`, decks in `tests/flux` and `tests/data/sector`",
        LEFT, TOP + Inches(4.8), FULL_W, size=15, align=PP_ALIGN.LEFT)

if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "sector-modelling.pptx"
    prs.save(out)
    print(f"Wrote {out} ({len(prs.slides)} slides)")
